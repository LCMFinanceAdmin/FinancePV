#!/usr/bin/env python3
"""Build a Word version of the handbook from the HTML one.

    python docs/handbook-to-docx.py

The HTML file is the original; this is a translation of it, and the two are
not the same document. Two things cannot survive the crossing and are handled
rather than lost:

  * The numbered rings over each screenshot are HTML elements positioned on top
    of the picture. Word has no equivalent, so they are drawn into the image
    itself here, using the same percentages the page uses. A figure arrives in
    Word as one flat picture with its labels already on it.

  * The cards, chips and coloured panels are CSS. They become Word's own
    constructs — shaded single-cell tables, real numbered lists, real tables —
    so the document can be edited by somebody who has never seen the HTML.

Headings are real Word heading styles, so the navigation pane works and the
table of contents can be built with a right-click.
"""
import io
import os
import pathlib
import re
import sys

try:
    from bs4 import BeautifulSoup, NavigableString, Tag
    from PIL import Image, ImageDraw, ImageFont
    import docx
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:
    sys.exit(f"Missing a dependency ({exc.name}). Try:\n"
             "    python -m pip install python-docx beautifulsoup4 pillow")

ROOT = pathlib.Path(__file__).resolve().parent
BOOK = ROOT / "lcm-finance-handbook.html"
IMGS = ROOT / "img"
OUT = ROOT / "LCM Finance Handbook.docx"
TMP = ROOT / ".docx-figures"

INK = RGBColor(0x1C, 0x27, 0x38)
BLUE = RGBColor(0x2F, 0x5B, 0x9C)
GREY = RGBColor(0x6B, 0x72, 0x80)
ORANGE = (194, 65, 12)

BODY_FONT = "Georgia"
HEAD_FONT = "Segoe UI Semibold"
MONO_FONT = "Consolas"

TEXT_WIDTH_CM = 16.4


# ── Word plumbing that python-docx does not expose ─────────────────────────
def shade(element, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    element.get_or_add_tcPr().append(shd) if element.__class__.__name__ == "_Cell" \
        else element._p.get_or_add_pPr().append(shd)


def cell_shade(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def left_bar(cell, hex_colour, size=18):
    """A thick left border — what carries a note box once its colour is gone
    on a black-and-white printer."""
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:color"), hex_colour)
    borders.append(left)
    cell._tc.get_or_add_tcPr().append(borders)


def dont_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))


def keep_with_next(par):
    pPr = par._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    pPr.append(el)


# ── Inline text ─────────────────────────────────────────────────────────────
def add_inline(par, node, bold=False, italic=False, mono=False, colour=None):
    for child in node.children:
        if isinstance(child, NavigableString):
            text = re.sub(r"\s+", " ", str(child))
            if not text:
                continue
            run = par.add_run(text)
            run.bold = bold
            run.italic = italic
            if mono:
                run.font.name = MONO_FONT
                run.font.size = Pt(9)
            if colour is not None:
                run.font.color.rgb = colour
        elif isinstance(child, Tag):
            if child.name == "br":
                par.add_run().add_break()
                continue
            cls = child.get("class", [])
            add_inline(
                par, child,
                bold or child.name in ("b", "strong"),
                italic or child.name in ("em", "i"),
                mono or "path" in cls,
                BLUE if "path" in cls else colour,
            )


def para(doc, node, size=10.5, space_after=6, colour=None, italic=False, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    add_inline(p, node, italic=italic, colour=colour)
    for run in p.runs:
        run.font.size = Pt(size)
        if colour is not None and run.font.color.rgb is None:
            run.font.color.rgb = colour
    return p


def text_para(doc, text, size=10.5, bold=False, italic=False, colour=None,
              space_after=6, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if colour is not None:
        run.font.color.rgb = colour
    return p


# ── Figures: draw the labels into the picture ───────────────────────────────
RING_RE = re.compile(r"--x:([\d.]+)%;--y:([\d.]+)%;--w:([\d.]+)%;--h:([\d.]+)%")
BADGE_RE = re.compile(r"--x:([\d.]+)%;--y:([\d.]+)%;")


def flatten_figure(fig, index):
    img = fig.find("img", attrs={"data-img": True})
    if img is None:
        return None
    name = img["data-img"]
    source = IMGS / name
    if not source.is_file():
        print(f"    {name} is not in docs/img — figure skipped")
        return None

    im = Image.open(source).convert("RGB")
    W, H = im.size
    draw = ImageDraw.Draw(im)

    for span in fig.select("span.anno-ring"):
        m = RING_RE.search(span.get("style", ""))
        if not m:
            continue
        x, y, w, h = (float(v) for v in m.groups())
        cx, cy, bw, bh = x / 100 * W, y / 100 * H, w / 100 * W, h / 100 * H
        draw.rounded_rectangle(
            [cx - bw / 2, cy - bh / 2, cx + bw / 2, cy + bh / 2],
            radius=max(8, W // 180), outline=ORANGE, width=max(4, W // 420))

    radius = max(16, W // 78)
    try:
        font = ImageFont.truetype("arialbd.ttf", int(radius * 1.25))
    except OSError:
        font = ImageFont.load_default()

    for span in fig.select("span.anno"):
        m = BADGE_RE.search(span.get("style", ""))
        if not m:
            continue
        x, y = (float(v) for v in m.groups())
        cx, cy = x / 100 * W, y / 100 * H
        draw.ellipse([cx - radius, cy - radius, cx + radius, cy + radius],
                     fill=ORANGE, outline=(255, 255, 255), width=max(3, W // 600))
        label = span.get_text(strip=True)
        box = draw.textbbox((0, 0), label, font=font)
        draw.text((cx - (box[2] - box[0]) / 2, cy - (box[3] - box[1]) / 2 - box[1]),
                  label, font=font, fill=(255, 255, 255))

    TMP.mkdir(exist_ok=True)
    out = TMP / f"fig{index:02d}-{name}"
    if im.width > 1200:
        im = im.resize((1200, round(im.height * 1200 / im.width)), Image.LANCZOS)
    im.save(out, optimize=True)
    # A phone screenshot at full text width would be a caricature of itself.
    narrow = "phone" in fig.get("class", []) or im.width < im.height
    return out, (Cm(7.5) if narrow else Cm(TEXT_WIDTH_CM))


def add_figure(doc, fig, index):
    made = flatten_figure(fig, index)
    if made is None:
        return
    path, width = made
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=width)


# ── Note boxes, cards and other panels ──────────────────────────────────────
def panel(doc, fill, bar, blocks, title=None):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    dont_split(table)
    cell = table.rows[0].cells[0]
    cell.width = Cm(TEXT_WIDTH_CM)
    cell_shade(cell, fill)
    left_bar(cell, bar)
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)

    if title:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x8A, 0x4B, 0x08) if bar != "2F5B9C" else BLUE
    for node in blocks:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        add_inline(p, node)
        for run in p.runs:
            run.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_note_box(doc, el):
    warn = "warn" in el.get("class", [])
    title_el = el.find("span", class_="t")
    title = title_el.get_text(" ", strip=True) if title_el else None
    panel(doc, "FDF6EC" if warn else "EEF4FD", "C2700C" if warn else "2F5B9C",
          el.find_all("p", recursive=False), title)


def add_steps(doc, ol):
    for i, li in enumerate(ol.find_all("li", recursive=False), 1):
        note = li.find("span", class_="note")
        if note:
            note.extract()
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, li)
        for run in p.runs:
            run.font.size = Pt(10.5)
        if note:
            q = doc.add_paragraph()
            q.paragraph_format.left_indent = Cm(0.9)
            q.paragraph_format.space_after = Pt(6)
            add_inline(q, note)
            for run in q.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = GREY


def add_numbered_key(doc, ol, marker="n"):
    """anno-key and callouts: a number in a circle, then the text."""
    for li in ol.find_all("li", recursive=False):
        badge = li.find("span", class_=marker)
        num = badge.get_text(strip=True) if badge else ""
        if badge:
            badge.extract()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(f"{num}   " if num else "")
        run.bold = True
        run.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C)
        run.font.size = Pt(10.5)
        add_inline(p, li)
        for r in p.runs[1:]:
            r.font.size = Pt(10.5)


def add_table(doc, tbl):
    head = tbl.find("thead")
    body = tbl.find("tbody")
    headers = [th.get_text(" ", strip=True) for th in head.find_all("th")] if head else []
    rows = body.find_all("tr") if body else tbl.find_all("tr")

    t = doc.add_table(rows=1 if headers else 0, cols=max(len(headers), 1))
    t.style = "Table Grid"
    if headers:
        for i, text in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            cell_shade(cell, "2F5B9C")

    for tr in rows:
        cells = tr.find_all(["th", "td"])
        row = t.add_row()
        for i, td in enumerate(cells):
            if i >= len(row.cells):
                break
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline(p, td)
            for run in p.runs:
                run.font.size = Pt(9)
                if td.name == "th":
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_stage_list(doc, items):
    """The spine: a stage name, then who it is waiting on."""
    for li in items:
        name = li.find(class_="stage-name")
        who = li.find(class_="stage-who")
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(1)
        if who and name and li.name == "div":
            # A tier states the step number first.
            run = p.add_run(who.get_text(" ", strip=True).upper() + "   ")
            run.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = BLUE
            run2 = p.add_run(name.get_text(" ", strip=True))
            run2.bold = True
            run2.font.size = Pt(10.5)
            body = li.find("p")
            if body:
                para(doc, body, size=10, space_after=7, indent=0.5)
            continue
        if name:
            run = p.add_run(name.get_text(" ", strip=True))
            run.bold = True
            run.font.size = Pt(10.5)
        if who:
            q = doc.add_paragraph()
            q.paragraph_format.left_indent = Cm(0.5)
            q.paragraph_format.space_after = Pt(6)
            add_inline(q, who)
            for run in q.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = GREY


def add_roles(doc, ul):
    for li in ul.find_all("li", recursive=False):
        name = li.find("b")
        text_para(doc, name.get_text(" ", strip=True) if name else "",
                  size=11, bold=True, colour=BLUE, space_after=2)
        inner = li.find("ul")
        if inner:
            for item in inner.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                add_inline(p, item)
                for run in p.runs:
                    run.font.size = Pt(10)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_branches(doc, div):
    for branch in div.find_all("div", class_="branch", recursive=False):
        cond = branch.find("span", class_="if")
        who = branch.find("b")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        keep_with_next(p)
        if cond:
            run = p.add_run(cond.get_text(" ", strip=True) + " ")
            run.font.size = Pt(9.5)
            run.font.color.rgb = GREY
        if who:
            run = p.add_run(who.get_text(" ", strip=True))
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = BLUE
        signers = branch.find("ol", class_="signers")
        if signers:
            for item in signers.find_all("li", recursive=False):
                q = doc.add_paragraph(style="List Bullet")
                q.paragraph_format.space_after = Pt(1)
                add_inline(q, item)
                for run in q.runs:
                    run.font.size = Pt(10)
        for body in branch.find_all("p", recursive=False):
            para(doc, body, size=9.5, space_after=8, colour=GREY)


def add_chips(doc, ul):
    for li in ul.find_all("li", recursive=False):
        d = li.find("span", class_="d")
        n = li.find("span", class_="n")
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(f"{d.get_text(strip=True)}  " if d else "")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = BLUE
        run2 = p.add_run(n.get_text(" ", strip=True) if n else "")
        run2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_scale(doc, div):
    for tier in div.find_all("div", class_="tier", recursive=False):
        amt = tier.find("span", class_="amt")
        sig = tier.find("span", class_="sig")
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(4)
        if amt:
            run = p.add_run(amt.get_text(" ", strip=True) + "   ")
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = BLUE
        if sig:
            run = p.add_run(sig.get_text(" ", strip=True))
            run.font.size = Pt(10.5)


def add_faq(doc, div):
    for det in div.find_all("details", recursive=False):
        summary = det.find("summary")
        if summary:
            p = text_para(doc, summary.get_text(" ", strip=True),
                          size=10.5, bold=True, colour=INK, space_after=2)
            keep_with_next(p)
        answer = det.find("div", class_="a")
        if answer:
            for body in answer.find_all("p", recursive=False) or [answer]:
                para(doc, body, size=10, space_after=8)


# ── The walk ────────────────────────────────────────────────────────────────
def render(doc, el, state):
    name = el.name
    cls = el.get("class", [])

    if name in ("h2", "h3", "h4"):
        level = {"h2": 1, "h3": 2, "h4": 3}[name]
        if level == 1:
            doc.add_page_break()
        h = doc.add_heading(el.get_text(" ", strip=True), level=level)
        for run in h.runs:
            run.font.name = HEAD_FONT
            run.font.color.rgb = BLUE if level == 1 else INK
            run.font.size = Pt({1: 17, 2: 12.5, 3: 11}[level])
        h.paragraph_format.space_before = Pt(4 if level == 1 else 12)
        h.paragraph_format.space_after = Pt(6)
        return

    if name == "p":
        if "figcap" in cls:
            p = para(doc, el, size=9, space_after=10, colour=GREY, italic=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif "lede" in cls:
            para(doc, el, size=12, space_after=10, colour=RGBColor(0x3A, 0x4A, 0x60))
        else:
            para(doc, el)
        return

    if name == "figure":
        state["figures"] += 1
        add_figure(doc, el, state["figures"])
        return

    if name == "div":
        if "note-box" in cls:
            add_note_box(doc, el); return
        if "tablewrap" in cls:
            for t in el.find_all("table", recursive=False):
                add_table(doc, t)
            return
        if "spine" in cls:
            add_stage_list(doc, el.find_all("div", class_="tier", recursive=False)); return
        if "branches" in cls:
            add_branches(doc, el); return
        if "scale" in cls:
            add_scale(doc, el); return
        if "faq" in cls:
            add_faq(doc, el); return
        for child in el.find_all(True, recursive=False):
            render(doc, child, state)
        return

    if name == "ol":
        if "steps" in cls:
            add_steps(doc, el); return
        if "anno-key" in cls:
            add_numbered_key(doc, el, "n"); return
        if "spine" in cls:
            add_stage_list(doc, el.find_all("li", recursive=False)); return
        for li in el.find_all("li", recursive=False):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, li)
            for run in p.runs:
                run.font.size = Pt(10.5)
        return

    if name == "ul":
        if "roles" in cls:
            add_roles(doc, el); return
        if "callouts" in cls:
            add_numbered_key(doc, el, "k"); return
        if "chips" in cls:
            add_chips(doc, el); return
        for li in el.find_all("li", recursive=False):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, li)
            for run in p.runs:
                run.font.size = Pt(10.5)
        return

    if name == "table":
        add_table(doc, el); return
    if name == "details":
        add_faq(doc, el.parent); return


def build_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rpr.set(qn("w:eastAsia"), BODY_FONT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-2" \h \z \u')
    inner = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Right-click here and choose Update Field to build the contents."
    inner.append(t)
    fld.append(inner)
    run._r.addnext(fld)


def main() -> int:
    if not BOOK.is_file():
        sys.exit(f"{BOOK} is not there.")
    soup = BeautifulSoup(BOOK.read_text(encoding="utf-8"), "html.parser")

    doc = Document()
    build_styles(doc)

    masthead = soup.find("header", class_="masthead")
    if masthead:
        eyebrow = masthead.find("p", class_="eyebrow")
        if eyebrow:
            text_para(doc, eyebrow.get_text(" ", strip=True).upper(), size=9.5,
                      bold=True, colour=BLUE, space_after=2)
        title = masthead.find("h1")
        if title:
            p = text_para(doc, title.get_text(" ", strip=True), size=26,
                          bold=True, colour=INK, space_after=8)
            for run in p.runs:
                run.font.name = HEAD_FONT
        stand = masthead.find("p", class_="standfirst")
        if stand:
            para(doc, stand, size=11.5, space_after=16, colour=RGBColor(0x3A, 0x4A, 0x60))

    text_para(doc, "Contents", size=13, bold=True, colour=BLUE, space_after=4)
    add_toc(doc)

    main_el = soup.find("main")
    state = {"figures": 0}
    for section in main_el.find_all("section", recursive=False):
        for child in section.find_all(True, recursive=False):
            render(doc, child, state)

    footer = soup.find("footer")
    if footer:
        doc.add_page_break()
        for body in footer.find_all("p"):
            para(doc, body, size=9, colour=GREY)

    doc.save(OUT)
    size = OUT.stat().st_size // 1024
    print(f"{state['figures']} figures flattened")
    print(f"Wrote {OUT.name} ({size} KB)")
    print("Open it and press Ctrl+A then F9 to fill in the contents page.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
